"""One-time asset builder. Runtime Publisher uses only the committed assets."""
from pathlib import Path
import shutil
import matplotlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor


def build():
    target = Path(__file__).resolve().parents[1] / 'publisher' / 'assets'
    fonts = target / 'fonts'
    fonts.mkdir(parents=True, exist_ok=True)
    source = Path(matplotlib.get_data_path()) / 'fonts' / 'ttf'
    for name in ('DejaVuSans.ttf', 'DejaVuSans-Bold.ttf', 'DejaVuSans-Oblique.ttf',
                 'DejaVuSans-BoldOblique.ttf', 'LICENSE_DEJAVU'):
        shutil.copyfile(source / name, fonts / name)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    normal = doc.styles['Normal']
    normal.font.name = 'DejaVu Sans'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(8)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[name].font.name = 'DejaVu Sans'
        doc.styles[name].font.color.rgb = RGBColor.from_string('1B2A4A')
    brand = doc.add_paragraph('BGTS INTERNATIONAL')
    brand.runs[0].font.color.rgb = RGBColor.from_string('C45C26')
    brand.runs[0].bold = True
    doc.add_paragraph('{{TITLE}}', style='Title')
    doc.add_paragraph('{{CLIENT}}')
    for key in ('CHALLENGE', 'APPROACH', 'TECHNOLOGY', 'OUTCOMES'):
        doc.add_paragraph('{{LABEL_' + key + '}}', style='Heading 1')
        doc.add_paragraph('{{' + key + '}}')
    section.footer.paragraphs[0].text = '{{CONFIDENTIAL}} — BGTS International'
    doc.save(target / 'case_study_template.docx')


if __name__ == '__main__':
    build()
