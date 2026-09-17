from docx import Document
from common.drafts import display_case_study
from publisher.assets import TEMPLATE
from publisher.publisher import render_docx
from tests.support import record, translated


def test_verified_literal_template_tokens_are_not_replaced(tmp_path):
    source=record()
    draft=display_case_study(translated(source,'en'),source['id'],'en')
    draft['sections']['challenge']='The source literally documents {{OUTCOMES}} and {{TITLE}}.'
    path=tmp_path/'literal.docx'
    render_docx(draft,TEMPLATE,path,source_record=source)
    text='\n'.join(p.text for p in Document(path).paragraphs)
    assert draft['sections']['challenge'] in text


def test_distinct_visible_citation_claim_is_rendered(tmp_path):
    source=record()
    draft=display_case_study(translated(source,'en'),source['id'],'en')
    draft['citations'][0]['claim']='The measured payment latency decreased.'
    path=tmp_path/'citation.docx'
    render_docx(draft,TEMPLATE,path,source_record=source)
    assert draft['citations'][0]['claim'] in '\n'.join(p.text for p in Document(path).paragraphs)
