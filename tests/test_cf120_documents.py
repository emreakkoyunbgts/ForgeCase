"""Contract evidence must see private text beyond a DOCX's top-level paragraphs."""
from io import BytesIO

from docx import Document
import pytest

from contract.documents import assert_document, document_text

PRIVATE = 'Synthetic Private Client'
EXPECTED = {'title': 'Ödemeler für Banken', 'outcome': 'İşlem gecikmesi %45 azaltıldı.'}


def document_bytes(story=None):
    document = Document()
    for value in EXPECTED.values():
        document.add_paragraph(value)
    if story is not None:
        if story == 'table':
            paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        elif story == 'nested-table':
            cell = document.add_table(rows=1, cols=1).cell(0, 0)
            paragraph = cell.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        elif story == 'header-table':
            paragraph = document.sections[0].header.add_table(rows=1, cols=1,
                width=document.sections[0].page_width).cell(0, 0).paragraphs[0]
        else:
            part = getattr(document.sections[0], story)
            paragraph = part.paragraphs[0]
        # A run boundary must not insert spaces that conceal the forbidden text.
        paragraph.add_run('Synthetic Pri')
        paragraph.add_run('vate Client')
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize('story', ['table', 'nested-table', 'header', 'first_page_header',
                                  'even_page_header', 'footer', 'first_page_footer',
                                  'even_page_footer', 'header-table'])
def test_private_text_in_secondary_document_story_cannot_pass_consent_assertion(story):
    content = document_bytes(story)
    text, pages = document_text(content, 'docx')
    assert PRIVATE in text and pages is None
    with pytest.raises(AssertionError, match='disclosed private text'):
        assert_document(content, 'docx', 'full-case-study', EXPECTED, forbidden=[PRIVATE])


def test_document_without_private_text_preserves_exact_unicode():
    text = assert_document(document_bytes(), 'docx', 'full-case-study', EXPECTED, forbidden=[PRIVATE])
    assert all(value in text for value in EXPECTED.values())


def test_unresolved_template_token_is_not_accepted_as_a_complete_document():
    document = Document(BytesIO(document_bytes()))
    document.sections[0].footer.paragraphs[0].text = '{{CLIENT_NAME}}'
    stream = BytesIO()
    document.save(stream)
    with pytest.raises(AssertionError, match='Unresolved document placeholder'):
        assert_document(stream.getvalue(), 'docx', 'full-case-study', EXPECTED)
